# Bygger data/Saldobalance_e-conomic_jan-aug_2026.xlsx: en tilfældig, e-conomic-agtig
# kontoplan for Nordhavn Composite A/S med beløb pr. måned (jan-aug 2026).
# Kvartalssummerne rammer præcis de realiserede tal i regnskabstabellen
# (ANNUAL_REPORT.q i financials.jsx), når kontiene mappes med standardmappingen.
# Brug (fra repo-roden): python data/make_trial_balance.py .
import json, os, random, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

ROOT = sys.argv[1] if len(sys.argv) > 1 else '.'
OUT = os.path.join(ROOT, 'data', 'Saldobalance_e-conomic_jan-aug_2026.xlsx')
rnd = random.Random(20260914)
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'Maj', 'Jun', 'Jul', 'Aug']
QM = [[0, 1, 2], [3, 4, 5], [6, 7]]          # månederne i Q1, Q2 og jul-aug
KR = lambda mio: int(round(mio * 1e6))

def split(total, shares, explicit=None):
    """Fordeler et beløb (kr) på konti efter andele; resten lander på kontoen med andel None."""
    explicit = explicit or {}
    out, rest_key = {}, None
    left = total - sum(explicit.values())
    base = left
    for k, s in shares:
        if s is None: rest_key = k; continue
        out[k] = int(round(base * s / 100 / 100) * 100)    # hele hundreder, som et bogført beløb ser ud
        left -= out[k]
    out[rest_key] = left
    out.update(explicit)
    return out

def months_of(q_amount, qi, profile):
    """Fordeler et kvartalsbeløb på dets måneder med lidt støj; sidste måned tager resten."""
    ms = QM[qi]
    w = {'rev': [[0.31, 0.32, 0.37], [0.33, 0.33, 0.34], [0.42, 0.58]],
         'flat': [[1/3] * 3, [1/3] * 3, [0.5, 0.5]]}[profile][qi]
    noise = [x * (1 + rnd.uniform(-0.07, 0.07)) for x in w]
    s = sum(noise)
    vals, used = [], 0
    for i, x in enumerate(noise[:-1]):
        v = int(round(q_amount * x / s))
        vals.append(v); used += v
    vals.append(q_amount - used)
    return dict(zip(ms, vals))

# ── Driftskonti: kvartalstal i tabellens fortegn (indtægt +, omkostning −), DKK mio. ──
Q = {
    'oms':   [10.60, 11.10, 7.38],
    'vare':  [-5.80, -6.08, -4.04],
    'pers':  [-3.58, -3.62, -2.43],
    'ext':   [-0.68, -0.70, -0.47],
    'afsk':  [-0.27, -0.27, -0.186],
    'fin':   [-0.11, -0.11, -0.08],
}
PL = {
    'oms': ('rev', [(1010, 'Salg af varer/ydelser m/moms', 29), (1020, 'Salg af varer/ydelser u/moms, EU', None),
                    (1030, 'Salg af varer/ydelser u/moms, verden', 15.5), (1040, 'Salg af serviceydelser m/moms', 3),
                    (1050, 'Lejeindtægter m/moms', 0)],
            {1090: ('Regulering igangværende arbejder', [0.04, -0.03, 0.02])}),
    'vare': ('rev', [(1310, 'Varekøb m/moms', 22), (1320, 'Varekøb u/moms, EU', None), (1330, 'Varekøb u/moms, import', 12),
                     (1340, 'Fremmed arbejde og underleverandører', 11), (1350, 'Fragt og told', 6), (1360, 'Emballage', 2)],
             {1380: ('Lagerregulering', [0.05, 0.10, 0.25])}),
    'pers': ('flat', [(2210, 'Lønninger', None), (2220, 'Feriepenge', 5), (2230, 'Pension', 8.5), (2240, 'ATP og AER', 1.2),
                      (2250, 'Sociale bidrag (AES, barsel.dk, FIB)', 1.2), (2270, 'Personalegoder og kantine', 1.8),
                      (2280, 'Uddannelse og kurser', 1)],
             {2290: ('Lønrefusioner', [0.02, 0.01, 0.015])}),
    'ext': ('flat', [(2710, 'Husleje', None), (2720, 'El, vand og varme', 16), (2730, 'Rengøring og vedligeholdelse af lokaler', 5),
                     (2810, 'Annoncer og reklame', 2), (2820, 'Messer og udstillinger', 3), (2830, 'Rejseudgifter', 5),
                     (2840, 'Restaurationsbesøg (repræsentation)', 1.2), (3010, 'Kontorartikler', 1), (3020, 'IT-licenser og software', 7),
                     (3030, 'Telefon og internet', 2), (3040, 'Revisor', 5), (3050, 'Advokat og rådgivning', 2), (3060, 'Forsikringer', 7),
                     (3070, 'Småanskaffelser', 2), (3080, 'Porto og gebyrer', 0.8), (3210, 'Brændstof', 3), (3220, 'Leasing af biler', 3),
                     (3230, 'Vedligeholdelse af biler', 0.8)], {}),
    'afsk': ('flat', [(3610, 'Afskrivning, udviklingsprojekter og software', 22.5), (3620, 'Afskrivning, bygninger', 12),
                      (3630, 'Afskrivning, produktionsanlæg og maskiner', None), (3640, 'Afskrivning, driftsmateriel og inventar', 10),
                      (3650, 'Avance/tab ved salg af anlægsaktiver', 0)], {}),
    'fin': ('flat', [(3820, 'Renteudgifter, bank og kassekredit', 42), (3830, 'Renteudgifter, realkredit', 25),
                     (3840, 'Renteudgifter, leasing', 10), (3850, 'Renter, lån fra ejer', 12), (3860, 'Kursdifferencer, valuta', None),
                     (3870, 'Bankgebyrer', 5)],
            {3810: ('Renteindtægter', [0.003, 0.002, 0.002])}),
}
pl_months = {}          # kontonr -> [8 beløb], tabellens fortegn, kr
pl_names = {}
for key, (profile, shares, explicit) in PL.items():
    for nr, name, _ in shares: pl_names[nr] = name
    for nr, (name, _) in explicit.items(): pl_names[nr] = name
    for qi in range(3):
        total = KR(Q[key][qi])
        ex = {nr: KR(v[qi]) for nr, (n, v) in explicit.items()}
        amounts = split(total, [(nr, s) for nr, _, s in shares], ex)
        for nr, a in amounts.items():
            m = months_of(a, qi, profile) if a else {i: 0 for i in QM[qi]}
            arr = pl_months.setdefault(nr, [0] * 8)
            for i, v in m.items(): arr[i] = v
for nr in (3410, 3420, 3910, 3920):
    pl_months[nr] = [0] * 8
pl_names.update({3410: 'Andre driftsindtægter', 3420: 'Andre driftsomkostninger', 3910: 'Skat af årets resultat', 3920: 'Regulering af udskudt skat'})

# ── Statuskonti: saldo primo og ultimo marts, juni og august (aktiver +, passiver +), DKK mio. ──
BS = [
    # nr, navn, side ('A' aktiv / 'P' passiv), primo, [mar, jun, aug]
    (5110, 'Udviklingsprojekter', 'A', 0.42, [0.43, 0.44, 0.44]),
    (5115, 'Software', 'A', 0.18, [0.17, 0.165, 0.16]),
    (5120, 'Grunde og bygninger', 'A', 2.40, [2.385, 2.37, 2.36]),
    (5130, 'Produktionsanlæg og maskiner', 'A', 2.50, [2.625, 2.73, 2.815]),
    (5140, 'Driftsmateriel og inventar', 'A', 0.40, [0.39, 0.395, 0.395]),
    (5150, 'Deposita', 'A', 0.10, [0.10, 0.10, 0.10]),
    (5210, 'Råvarer og hjælpematerialer', 'A', 1.60, [1.65, 1.70, 1.80]),
    (5220, 'Varer under fremstilling', 'A', 0.70, [0.72, 0.75, 0.85]),
    (5230, 'Færdigvarer', 'A', 0.90, [0.88, 0.90, 0.95]),
    (5300, 'Debitorer', 'A', 2.55, [2.63, 2.70, 2.90]),
    (5310, 'Andre tilgodehavender', 'A', 0.20, [0.22, 0.21, 0.22]),
    (5320, 'Forudbetalte omkostninger', 'A', 0.15, [0.16, 0.17, 0.17]),
    (5400, 'Bank', 'A', 1.89, [1.99, 2.14, 2.07]),
    (5410, 'Kasse', 'A', 0.01, [0.01, 0.01, 0.01]),
    (6110, 'Selskabskapital', 'P', 1.00, [1.00, 1.00, 1.00]),
    (6120, 'Overført resultat', 'P', 5.20, [5.20, 5.20, 5.20]),
    (6150, 'Udskudt skat', 'P', 0.00, [0.00, 0.00, 0.00]),
    (6210, 'Realkreditlån', 'P', 1.90, [1.87, 1.84, 1.82]),
    (6220, 'Banklån', 'P', 1.55, [1.50, 1.45, 1.42]),
    (6230, 'Leasingforpligtelser', 'P', 0.65, [0.63, 0.71, 0.71]),
    (6240, 'Lån fra ejer', 'P', 0.50, [0.50, 0.50, 0.50]),
    (6310, 'Kassekredit', 'P', 0.45, [0.52, 0.48, 0.60]),
    (6320, 'Kortfristet del af langfristet gæld', 'P', 0.45, [0.45, 0.45, 0.45]),
    (6330, 'Leverandører af varer og tjenesteydelser', 'P', 1.30, [1.42, 1.50, 1.62]),
    (6340, 'Forudbetalinger fra kunder', 'P', 0.45, [0.40, 0.42, 0.45]),
    (6350, 'Skyldig A-skat og AM-bidrag', 'P', 0.12, [0.13, 0.13, 0.13]),
    (6360, 'Skyldige feriepenge', 'P', 0.18, [0.26, 0.30, 0.326]),
    (6370, 'Skyldig moms', 'P', 0.15, [0.20, 0.19, 0.22]),
    (6380, 'Skyldig selskabsskat', 'P', 0.05, [0.05, 0.05, 0.05]),
    (6390, 'Anden gæld', 'P', 0.05, [0.07, 0.08, 0.09]),
]
QEND = [2, 5, 7]
ytd_profit = []          # kr, ultimo hver måned
acc = 0
for m in range(8):
    acc += sum(v[m] for v in pl_months.values())
    ytd_profit.append(acc)
# Saldo ultimo hver måned i e-conomic-fortegn (debet +). Banken udligner, så
# saldobalancen går i nul hver måned; i kvartalernes sidste måned er den præcis målet.
bal = {}
for nr, name, side, primo, ends in BS:
    sg = 1 if side == 'A' else -1
    p = KR(primo); pts = [p] + [KR(e) for e in ends]
    s = []
    for m in range(8):
        if m in QEND: s.append(pts[QEND.index(m) + 1]); continue
        qi = 0 if m < 3 else 1 if m < 6 else 2
        a, b = pts[qi], pts[qi + 1]
        start = -1 if qi == 0 else QEND[qi - 1]
        f = (m - start) / (QEND[qi] - start)
        v = a + (b - a) * f
        if nr not in (6110, 6120, 6150, 5150, 6320, 6380, 6240):
            v *= 1 + rnd.uniform(-0.02, 0.02)
        s.append(int(round(v / 100) * 100))
    bal[nr] = {'name': name, 'primo': sg * p, 'ends': [sg * x for x in s]}
for m in range(8):
    # sum af alle statuskonti + periodens resultat (driftskonti i e-conomic-fortegn) = 0
    others = sum(b['ends'][m] for nr, b in bal.items() if nr != 5400)
    pl_erp = -ytd_profit[m]
    bal[5400]['ends'][m] = -(others + pl_erp)
for i, m in enumerate(QEND):
    assert bal[5400]['ends'][m] == KR([1.99, 2.14, 2.07][i]), ('bank', m, bal[5400]['ends'][m])

# ── Kontoplanen i e-conomic-rækkefølge ──
S, H = 'Sum', 'Overskrift'
plan = [
    (1000, 'RESULTATOPGØRELSE', H), (1001, 'Omsætning', H),
    *[(n, None, 'Drift') for n in (1010, 1020, 1030, 1040, 1050, 1090)],
    (1099, 'Omsætning i alt', S, 1001, 1098),
    (1300, 'Direkte omkostninger', H),
    *[(n, None, 'Drift') for n in (1310, 1320, 1330, 1340, 1350, 1360, 1380)],
    (1399, 'Direkte omkostninger i alt', S, 1300, 1398),
    (1400, 'DÆKNINGSBIDRAG', S, 1000, 1399),
    (2200, 'Personaleomkostninger', H),
    *[(n, None, 'Drift') for n in (2210, 2220, 2230, 2240, 2250, 2270, 2280, 2290)],
    (2299, 'Personaleomkostninger i alt', S, 2200, 2298),
    (2700, 'Lokaleomkostninger', H), *[(n, None, 'Drift') for n in (2710, 2720, 2730)],
    (2799, 'Lokaleomkostninger i alt', S, 2700, 2798),
    (2800, 'Salgs- og rejseomkostninger', H), *[(n, None, 'Drift') for n in (2810, 2820, 2830, 2840)],
    (2899, 'Salgs- og rejseomkostninger i alt', S, 2800, 2898),
    (3000, 'Administrationsomkostninger', H), *[(n, None, 'Drift') for n in (3010, 3020, 3030, 3040, 3050, 3060, 3070, 3080)],
    (3099, 'Administrationsomkostninger i alt', S, 3000, 3098),
    (3200, 'Bilomkostninger', H), *[(n, None, 'Drift') for n in (3210, 3220, 3230)],
    (3299, 'Bilomkostninger i alt', S, 3200, 3298),
    (3400, 'Andre driftsposter', H), *[(n, None, 'Drift') for n in (3410, 3420)],
    (3500, 'RESULTAT FØR AFSKRIVNINGER', S, 1000, 3499),
    (3600, 'Afskrivninger', H), *[(n, None, 'Drift') for n in (3610, 3620, 3630, 3640, 3650)],
    (3699, 'Afskrivninger i alt', S, 3600, 3698),
    (3700, 'RESULTAT FØR RENTER', S, 1000, 3699),
    (3800, 'Finansielle poster', H), *[(n, None, 'Drift') for n in (3810, 3820, 3830, 3840, 3850, 3860, 3870)],
    (3899, 'Finansielle poster i alt', S, 3800, 3898),
    (3900, 'RESULTAT FØR SKAT', S, 1000, 3899),
    *[(n, None, 'Drift') for n in (3910, 3920)],
    (3999, 'ÅRETS RESULTAT', S, 1000, 3998),
    (5000, 'BALANCE', H), (5100, 'Anlægsaktiver', H),
    *[(n, None, 'Status') for n in (5110, 5115, 5120, 5130, 5140, 5150)],
    (5199, 'Anlægsaktiver i alt', S, 5100, 5198),
    (5200, 'Omsætningsaktiver', H),
    *[(n, None, 'Status') for n in (5210, 5220, 5230, 5300, 5310, 5320, 5400, 5410)],
    (5499, 'Omsætningsaktiver i alt', S, 5200, 5498),
    (5500, 'AKTIVER I ALT', S, 5000, 5499),
    (6000, 'PASSIVER', H), (6100, 'Egenkapital', H), *[(n, None, 'Status') for n in (6110, 6120)],
    (6130, 'Egenkapital i alt', S, 6100, 6129),
    (6140, 'Hensatte forpligtelser', H), (6150, None, 'Status'),
    (6200, 'Langfristet gæld', H), *[(n, None, 'Status') for n in (6210, 6220, 6230, 6240)],
    (6299, 'Langfristet gæld i alt', S, 6200, 6298),
    (6300, 'Kortfristet gæld', H), *[(n, None, 'Status') for n in (6310, 6320, 6330, 6340, 6350, 6360, 6370, 6380, 6390)],
    (6399, 'Kortfristet gæld i alt', S, 6300, 6398),
    (6500, 'PASSIVER I ALT', S, 6000, 6499),
]

# ── Excel ──
wb = Workbook()
ws = wb.active
ws.title = 'Saldobalance'
F = lambda size=10, **k: Font(name='Arial', size=size, **k)
INPUT = F(color='0000FF')
thin = Side(style='thin', color='BFBFBF')
ws['A1'] = 'Saldobalance pr. måned'; ws['A1'].font = F(bold=True, size=13)
ws['A2'] = 'Nordhavn Composite A/S · CVR 38427156'; ws['A2'].font = F()
ws['A3'] = 'Kilde: e-conomic, eksporteret 14-09-2026 · Periode: januar-august 2026 · Beløb i DKK'; ws['A3'].font = F(color='595959')
ws['A4'] = ('Fortegn som i e-conomic: debet er plus, kredit er minus. Driftskonti: månedens bevægelse. '
            'Statuskonti: saldo primo og månedens bevægelse. Blå tal er indtastede, sorte er beregnede.'); ws['A4'].font = F(color='595959')
HDR = 6
heads = ['Kontonr', 'Kontonavn', 'Kontotype', 'Sum fra', 'Sum til', 'Primo 01-01-2026'] + [m + ' 2026' for m in MONTHS] + ['Saldo 31-08-2026']
for c, h in enumerate(heads, 1):
    cell = ws.cell(row=HDR, column=c, value=h)
    cell.font = F(bold=True); cell.fill = PatternFill('solid', fgColor='F2F2F2')
    cell.border = Border(bottom=thin)
    cell.alignment = Alignment(horizontal='right' if c >= 4 else 'left', vertical='bottom', wrap_text=True)
first = HDR + 1
last = first + len(plan) - 1
NUM = '#,##0;-#,##0;"-"'
COL = lambda c: ws.cell(row=1, column=c).column_letter
for i, row in enumerate(plan):
    r = first + i
    nr, name, typ = row[0], row[1], row[2]
    if name is None: name = pl_names.get(nr) or bal[nr]['name']
    ws.cell(row=r, column=1, value=nr).font = F(bold=typ in (H, S))
    ws.cell(row=r, column=2, value=name).font = F(bold=typ in (H, S))
    ws.cell(row=r, column=3, value=typ).font = F(color='595959')
    if typ == S:
        ws.cell(row=r, column=4, value=row[3]).font = F(color='595959')
        ws.cell(row=r, column=5, value=row[4]).font = F(color='595959')
        for c in range(6, 6 + 9):
            L = COL(c)
            # Området slutter på rækken over sumkontoen, så formlen ikke peger på sig selv
            # (intervallets konti står altid over sumkontoen)
            e = r - 1
            ws.cell(row=r, column=c, value=f'=SUMIFS({L}${first}:{L}{e},$A${first}:$A{e},">="&$D{r},$A${first}:$A{e},"<="&$E{r},$C${first}:$C{e},"<>Sum")').font = F(bold=True)
    elif typ in ('Drift', 'Status'):
        if typ == 'Drift':
            vals = [-v for v in pl_months[nr]]               # e-conomic: indtægt kredit (−), omkostning debet (+)
            ws.cell(row=r, column=6, value=None)
        else:
            b = bal[nr]
            ends = b['ends']; prev = b['primo']; vals = []
            for e in ends: vals.append(e - prev); prev = e
            ws.cell(row=r, column=6, value=b['primo']).font = INPUT
        for m, v in enumerate(vals):
            ws.cell(row=r, column=7 + m, value=v).font = INPUT
    if typ in ('Drift', 'Status', S):
        ws.cell(row=r, column=15, value=f'=SUM(F{r}:N{r})').font = F(bold=typ == S)
    for c in range(6, 16): ws.cell(row=r, column=c).number_format = NUM
    if typ == S:
        for c in range(1, 16): ws.cell(row=r, column=c).border = Border(top=thin)
# Kontrolrække: alle drifts- og statuskonti tilsammen skal give 0 i hver kolonne
cr = last + 2
ws.cell(row=cr, column=2, value='Kontrol: debet minus kredit (skal være 0)').font = F(bold=True)
for c in range(6, 16):
    L = COL(c)
    ws.cell(row=cr, column=c, value=f'=SUMIFS({L}${first}:{L}${last},$C${first}:$C${last},"Drift")+SUMIFS({L}${first}:{L}${last},$C${first}:$C${last},"Status")').font = F(bold=True)
    ws.cell(row=cr, column=c).number_format = NUM
ws.cell(row=HDR, column=1).comment = Comment('Prototypen læser Kontonr, Kontonavn, Kontotype, Primo og månedskolonnerne. Sumkonti beregnes.', 'Crediwire')
widths = [9, 42, 11, 8, 8, 14] + [12] * 8 + [15]
for c, w in enumerate(widths, 1): ws.column_dimensions[COL(c)].width = w
ws.freeze_panes = ws.cell(row=first, column=3)
ws.row_dimensions[HDR].height = 28

# Læs mig
rd = wb.create_sheet('Læs mig')
lines = [
    ('Saldobalance pr. måned fra e-conomic', True),
    ('', False),
    ('Arket er kilden til de realiserede perioder (Q1, Q2 og jul-aug 2026) i Regnskab under Virksomheden.', False),
    ('Prototypen henter filen, mapper hver konto til en Crediwire-kategori og lægger beløbene sammen pr. kategori.', False),
    ('', False),
    ('Kolonner', True),
    ('Kontonr og Kontonavn: kontoplanen, som den står i e-conomic.', False),
    ('Kontotype: Overskrift, Drift (resultatopgørelse), Status (balance) eller Sum.', False),
    ('Sum fra / Sum til: kontointervallet, en sumkonto lægger sammen.', False),
    ('Primo 01-01-2026: statuskontiens saldo ved årets start (driftskonti er tomme).', False),
    ('Jan-Aug 2026: månedens bevægelse på kontoen.', False),
    ('Saldo 31-08-2026: primo plus bevægelserne (driftskonti: år til dato).', False),
    ('', False),
    ('Fortegn', True),
    ('Som i e-conomic: debet er plus, kredit er minus. Omsætning, gæld og egenkapital er derfor negative.', False),
    ('', False),
    ('Sådan retter du', True),
    ('Ret kun de blå tal. Sumkonti, saldo og kontrolrækken nederst er formler.', False),
    ('Kontrolrækken skal give 0 i alle kolonner. Gør den ikke det, går debet og kredit ikke op.', False),
    ('Gem filen samme sted og genindlæs prototypen. Nye konti kommer med som "Ikke mappet", til de mappes.', False),
]
for i, (txt, bold) in enumerate(lines, 1):
    rd.cell(row=i, column=1, value=txt).font = F(bold=bold, size=12 if i == 1 else 10)
rd.column_dimensions['A'].width = 110
wb.calculation.fullCalcOnLoad = True
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

# Kontrol i Python: kvartalssummer pr. gruppe og saldobalancen i nul
def qsum(nrs):
    return [sum(pl_months[n][m] for n in nrs for m in QM[q]) / 1e6 for q in range(3)]
groups = {k: [nr for nr, _, _ in v[1]] + list(v[2].keys()) for k, v in PL.items()}
check = {k: [round(x, 6) for x in qsum(n)] for k, n in groups.items()}
mov_ok = all(sum(-pl_months[n][m] for n in pl_months) + sum((b['ends'][m] - (b['ends'][m - 1] if m else b['primo'])) for b in bal.values()) == 0 for m in range(8))
primo_ok = sum(b['primo'] for b in bal.values()) == 0
print(json.dumps({'file': OUT, 'rows': len(plan), 'quarters': check, 'movements_balance': mov_ok, 'primo_balance': primo_ok,
                  'bank_month_end': [round(x / 1e6, 3) for x in bal[5400]['ends']]}, ensure_ascii=False, indent=1))
